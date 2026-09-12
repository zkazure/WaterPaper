"""docx_io 回归测试。

不做全面覆盖，只钉住实际踩到过的两处缺陷，保证它们不会复活：

1. `_add_markdown_runs` 用了 `Pt` / `qn` 却没导入 → 只要走模板（`props` 非空）就
   `NameError`。这条路径正是「检测/改写模式」保留原格式写回的主路径。
2. 正文级别 fallback 的排序 key 只有 `-count`，与注释声称的「最小字号组」不符 ——
   并列时稳定排序保留的是更大字号，正文被套成上一级标题的字号。
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

import docx_io
import pytest


def _make_template(path: Path, entries: Sequence[tuple[str, float]]) -> Path:
    """造一个各级字号分明的「学校模板」。

    entries 形如 [("标题", 22.0), ("一级标题", 16.0), ...]，字号种类决定分级。
    """
    docx = pytest.importorskip("docx", reason="需要 python-docx")
    from docx.shared import Pt

    doc = docx.Document()
    for text, size in entries:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.font.name = "宋体"
    doc.save(str(path))
    return path


# 22/16/14/12 四档：title / heading1 / heading2 / heading3(与正文同档触发 fallback)
FOUR_LEVEL: list[tuple[str, float]] = [
    ("模板标题", 22.0),
    ("一级标题示例", 16.0),
    ("二级标题示例", 14.0),
    ("正文示例文字", 12.0),
]


def test_add_markdown_runs_applies_size_without_nameerror():
    """回归（缺陷 1）：`Pt` 未导入 → NameError: name 'Pt' is not defined。

    该块紧跟在每次 `_add_markdown_runs(...)` 调用之后还有一次 `_apply_format(...)`，
    所以缺导入时不会静默降级，而是直接抛异常中断整篇生成。
    """
    docx = pytest.importorskip("docx", reason="需要 python-docx")
    doc = docx.Document()
    para = doc.add_paragraph()
    props = {"size_pt": 14, "eastAsia": "宋体", "ascii": "Times New Roman"}

    docx_io._add_markdown_runs(para, "正文**加粗**文字", props, {})

    assert para.runs[-1].font.size.pt == 14  # 不再 NameError
    run_fonts = para.runs[-1]._element.rPr.rFonts
    assert run_fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") == "宋体"


def test_add_markdown_runs_splits_bold_and_italic():
    """顺带钉住 run 切分：字号块不能把 ** 标记弄丢或漏进正文。"""
    docx = pytest.importorskip("docx", reason="需要 python-docx")
    doc = docx.Document()
    para = doc.add_paragraph()

    docx_io._add_markdown_runs(para, "前缀**加粗**后缀", {}, {})

    assert [r.text for r in para.runs] == ["前缀", "加粗", "后缀"]
    assert para.runs[1].font.bold is True
    assert "*" not in "".join(r.text for r in para.runs)


def test_extract_template_formats_body_prefers_smaller_size_on_tie(tmp_path):
    """回归（缺陷 2）：各级段落数并列时，正文必须取更小字号。

    注释与意图都是「the smallest size group with the most paragraphs」，但原实现的
    key 只有 `-count`；并列时 Python 稳定排序保留的是输入顺序（字号降序），
    于是正文套上了二级标题的 14pt 而不是 12pt。
    """
    docx = pytest.importorskip("docx", reason="需要 python-docx")
    template = _make_template(tmp_path / "t.docx", FOUR_LEVEL)

    fmt = docx_io._extract_template_formats(docx.Document(str(template)))

    assert fmt["body"]["size_pt"] == 12.0
    assert fmt["heading1"]["size_pt"] == 16.0
    assert fmt["heading2"]["size_pt"] == 14.0


def test_formatted_write_applies_template_sizes_end_to_end(tmp_path):
    """产物级验证：跑完整 formatted_write，逐段核对写出的字号。

    这是两个缺陷共同的验收点 —— 缺陷 1 会让它直接崩，缺陷 2 会让正文段字号偏大。
    """
    docx = pytest.importorskip("docx", reason="需要 python-docx")
    template = _make_template(tmp_path / "template.docx", FOUR_LEVEL)
    out = tmp_path / "out.docx"

    docx_io.formatted_write_docx(
        str(out), "# 一级标题\n\n## 二级标题\n\n正文**加粗**内容。\n", str(template))

    sizes = {p.text: p.runs[0].font.size.pt
             for p in docx.Document(str(out)).paragraphs if p.text.strip()}
    assert sizes["一级标题"] == 16.0
    assert sizes["二级标题"] == 14.0
    assert sizes["正文加粗内容。"] == 12.0


def test_formatted_write_without_template_does_not_crash(tmp_path):
    """无模板（props 为空）时不应进字号分支 —— 确认修复没有把这条路弄坏。"""
    docx = pytest.importorskip("docx", reason="需要 python-docx")
    out = tmp_path / "plain.docx"

    docx_io.formatted_write_docx(str(out), "# 标题\n\n正文。\n")

    assert [p.text for p in docx.Document(str(out)).paragraphs if p.text.strip()] \
        == ["标题", "正文。"]
